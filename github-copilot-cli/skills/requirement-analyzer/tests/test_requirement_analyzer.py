from __future__ import annotations

import importlib.util
import argparse
import contextlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import warnings
import zipfile
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPT = SKILL_ROOT / "scripts" / "requirement_analyzer.py"
FIXTURE = SKILL_ROOT / "tests" / "fixtures" / "basic"

spec = importlib.util.spec_from_file_location("requirement_analyzer", SCRIPT)
assert spec and spec.loader
analyzer = importlib.util.module_from_spec(spec)
spec.loader.exec_module(analyzer)


def project_directory():
    return tempfile.TemporaryDirectory(dir=SKILL_ROOT / "tests")


class RequirementAnalyzerTests(unittest.TestCase):
    maxDiff = None

    @staticmethod
    def call_in_process(function, **arguments):
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            result = function(argparse.Namespace(**arguments))
        if result:
            raise AssertionError(output.getvalue())
        return json.loads(output.getvalue())

    @staticmethod
    def read_run(prepared):
        run = json.loads(Path(prepared["run"]).read_text(encoding="utf-8"))
        manifest = json.loads(Path(prepared["manifest"]).read_text(encoding="utf-8"))
        return run, manifest

    def publish_draft(self, run, manifest):
        draft = self.make_draft_ledger(run, manifest)
        path = Path(run["run_directory"]) / "completed.json"
        path.write_text(json.dumps(draft), encoding="utf-8")
        return self.call_in_process(analyzer._publish, run=str(path.parent / "run.json"), ledger=str(path))

    def run_cli(self, *arguments: str, expected: int = 0) -> dict:
        completed = subprocess.run(
            [sys.executable, str(SCRIPT), *arguments],
            capture_output=True,
            text=True,
            check=False,
        )
        self.assertEqual(
            expected,
            completed.returncode,
            msg=f"stdout:\n{completed.stdout}\nstderr:\n{completed.stderr}",
        )
        payload = completed.stdout.strip() or completed.stderr.strip()
        return json.loads(payload)

    def prepare_fixture(self, temporary: Path) -> tuple[dict, dict]:
        source_base = temporary / "source"
        shutil.copytree(FIXTURE, source_base)
        legacy_requirements = source_base / "Requirements"
        requirements = source_base / "requirements"
        legacy_requirements.rename(requirements)
        config = source_base / "lisa-config.json"
        config.write_text(json.dumps({"basePath": "."}), encoding="utf-8")

        message = EmailMessage()
        message["Subject"] = "Requirement example"
        message.set_content(
            "The assistant should preserve source traceability. "
            "Policy reference: https://example.com/policy"
        )
        message.add_attachment(
            b"Embedded requirement evidence.",
            maintype="text",
            subtype="plain",
            filename="attachment.txt",
        )
        analyzer._atomic_write_text(requirements / "example.eml", message.as_string())
        os.utime(requirements / "example.eml", (946684800, 946684800))
        (requirements / "example.eml").chmod(0o444)

        output = source_base / "output"
        prepared = self.run_cli(
            "prepare",
            "--config",
            str(config),
            "--workers",
            "2",
        )
        run = json.loads(Path(prepared["run"]).read_text(encoding="utf-8"))
        manifest = json.loads(Path(prepared["manifest"]).read_text(encoding="utf-8"))
        self.assertTrue(Path(prepared["review_pack"]).exists())
        analysis = output / "analysis"
        self.assertEqual(output.resolve(), Path(run["temp_output_path"]))
        self.assertEqual(analysis.resolve(), Path(run["output_root"]))
        self.assertEqual(["analysis"], sorted(item.name for item in output.iterdir()))
        self.assertTrue(Path(prepared["run"]).is_relative_to(analysis))
        self.assertEqual(analysis.resolve(), Path(prepared["target_markdown"]).parent)
        return run, manifest

    @staticmethod
    def make_draft_ledger(run: dict, manifest: dict) -> dict:
        findings = []
        source_annotations = []
        next_number = 1

        def add_finding(
            kind: str,
            statement: str,
            status: str,
            confidence: str,
            source_id: str,
            locator: str,
            evidence_type: str,
        ) -> str:
            nonlocal next_number
            identifier = f"D-{next_number:03d}"
            next_number += 1
            findings.append(
                {
                    "finding_id": identifier,
                    "kind": kind,
                    "statement": statement,
                    "status": status,
                    "confidence": confidence,
                    "evidence": [
                        {
                            "source_id": source_id,
                            "locator": locator,
                            "evidence_type": evidence_type,
                        }
                    ],
                }
            )
            return identifier

        for source in manifest["sources"]:
            identifier = add_finding(
                "Observed fact",
                f"The source inventory contains {source['relative_path']}.",
                "Current",
                "Confirmed",
                source["source_id"],
                "filesystem metadata",
                "metadata",
            )
            source_annotations.append(
                {
                    "source_id": source["source_id"],
                    "role": "Requirement evidence",
                    "classification": "Folder file",
                    "finding_ids": [identifier],
                }
            )

        first_source = manifest["sources"][0]["source_id"]
        conversational = add_finding(
            "Derived classification",
            "The requested assistant interaction is conversational.",
            "Confirmed",
            "Confirmed",
            first_source,
            "line 1",
            "derived",
        )
        platform_gap = add_finding(
            "Analyst-identified gap",
            "No permitted agent-development platform is evidenced.",
            "Not evidenced",
            "Missing",
            "CORPUS",
            "platform completeness check across the extracted corpus",
            "absence_check",
        )
        knowledge_gap = add_finding(
            "Analyst-identified gap",
            "No source qualifies as a runtime knowledge source.",
            "Not evidenced",
            "Missing",
            "CORPUS",
            "knowledge-source qualification check across the extracted corpus",
            "absence_check",
        )
        integration_gap = add_finding(
            "Analyst-identified gap",
            "No integration is evidenced.",
            "Not evidenced",
            "Missing",
            "CORPUS",
            "integration completeness check across the extracted corpus",
            "absence_check",
        )
        autonomous_gap = add_finding(
            "Analyst-identified gap",
            "Autonomous behavior is not evidenced.",
            "Not evidenced",
            "Missing",
            "CORPUS",
            "agentic-behavior completeness check across the extracted corpus",
            "absence_check",
        )
        multi_agent_gap = add_finding(
            "Analyst-identified gap",
            "Child-agent or multi-agent behavior is not evidenced.",
            "Not evidenced",
            "Missing",
            "CORPUS",
            "agentic-behavior completeness check across the extracted corpus",
            "absence_check",
        )
        oversight_gap = add_finding(
            "Analyst-identified gap",
            "Human handoff or oversight is not evidenced.",
            "Not evidenced",
            "Missing",
            "CORPUS",
            "agentic-behavior completeness check across the extracted corpus",
            "absence_check",
        )

        first_inventory_finding = source_annotations[0]["finding_ids"][0]
        section_names = [
            "Executive Summary",
            "Problem Statement",
            "Current State",
            "Desired Future State",
            "Goals",
            "Success Criteria",
            "Metrics and Baselines",
            "Data Sources",
            "Data Types",
            "Dependencies and Constraints",
            "Solution Components",
            "Scope and Delivery Phases",
            "Gaps and Conflicts",
        ]
        sections = {
            name: [
                {
                    "type": "paragraph",
                    "text": f"{name} is grounded in the inventoried requirement evidence.",
                    "finding_ids": [first_inventory_finding],
                }
            ]
            for name in section_names
        }

        return {
            "schema_version": "1.0",
            "run_id": run["run_id"],
            "findings": findings,
            "source_annotations": source_annotations,
            "referred_artifacts": [],
            "manual_reviews": [],
            "batch_reviews": analyzer._batch_reviews(
                json.loads(Path(run["review_index_path"]).read_text(encoding="utf-8"))
            ),
            "knowledge_sources": [],
            "knowledge_source_notes": [
                {
                    "text": (
                        "No confirmed or strong-candidate runtime knowledge sources "
                        "were evidenced."
                    ),
                    "finding_ids": [knowledge_gap],
                }
            ],
            "platforms": [],
            "platform_absence_finding_ids": [platform_gap],
            "integrations": [],
            "integration_notes": [
                {
                    "text": "No integrations were evidenced.",
                    "finding_ids": [integration_gap],
                }
            ],
            "agentic_behaviors": [
                {
                    "behavior": "Conversational",
                    "requirement_status": "Confirmed",
                    "evidenced_behavior": "Users ask the assistant questions.",
                    "trigger_decision_handoff": "User prompt",
                    "gaps": "Channel and response targets are unspecified.",
                    "finding_ids": [conversational],
                },
                {
                    "behavior": "Autonomous",
                    "requirement_status": "Not evidenced",
                    "evidenced_behavior": "No autonomous behavior is described.",
                    "trigger_decision_handoff": "Not evidenced",
                    "gaps": "No autonomous requirements are supplied.",
                    "finding_ids": [autonomous_gap],
                },
                {
                    "behavior": "Child-Agent/Multi-Agent",
                    "requirement_status": "Not evidenced",
                    "evidenced_behavior": "No delegation or multi-agent behavior is described.",
                    "trigger_decision_handoff": "Not evidenced",
                    "gaps": "No multi-agent requirements are supplied.",
                    "finding_ids": [multi_agent_gap],
                },
                {
                    "behavior": "Human Handoff/oversight",
                    "requirement_status": "Not evidenced",
                    "evidenced_behavior": "No human handoff is described.",
                    "trigger_decision_handoff": "Not evidenced",
                    "gaps": "No oversight requirement is supplied.",
                    "finding_ids": [oversight_gap],
                },
            ],
            "sections": sections,
        }

    def normalize_draft(
        self, run: dict, draft: dict, name: str = "test"
    ) -> Path:
        draft_path = Path(run["run_directory"]) / f"ledger.{name}.draft.json"
        draft_path.write_text(json.dumps(draft, indent=2), encoding="utf-8")
        normalized_path = (
            Path(run["run_directory"]) / f"ledger.{name}.normalized.json"
        )
        self.run_cli(
            "normalize",
            "--run",
            str(Path(run["run_directory"]) / "run.json"),
            "--ledger",
            str(draft_path),
            "--output",
            str(normalized_path),
        )
        return normalized_path

    def test_resolves_only_exact_configured_root(self) -> None:
        with project_directory() as directory:
            base = Path(directory)
            exact = base / "requirements"
            exact.mkdir()
            self.assertEqual(exact.resolve(), analyzer.resolve_requirements_root(exact))

        with project_directory() as directory:
            base = Path(directory)
            with self.assertRaises(analyzer.AnalyzerError):
                analyzer.resolve_requirements_root(base)

    def test_native_office_and_email_extraction(self) -> None:
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation

        docx = io.BytesIO()
        document = Document()
        document.add_paragraph("Document requirement")
        document.save(docx)
        result = analyzer._extract_bytes(
            "requirement.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            docx.getvalue(),
        )
        self.assertEqual("manual-review-required", result["status"])
        self.assertEqual("render:all-pages", result["review_targets"][0]["target_key"])
        self.assertIn("Document requirement", analyzer._collect_extracted_text(result))

        xlsx = io.BytesIO()
        workbook = Workbook()
        workbook.active["A1"] = "Spreadsheet requirement"
        workbook.save(xlsx)
        result = analyzer._extract_bytes(
            "requirement.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            xlsx.getvalue(),
        )
        self.assertEqual("complete", result["status"])
        self.assertEqual("xlsx-stream-lossless", result["method"])
        self.assertEqual(0, len(result["review_targets"]))
        self.assertEqual(1, result["metadata"]["rows_scanned"])
        self.assertIn("Spreadsheet requirement", analyzer._collect_extracted_text(result))
        row_number, values = analyzer._parse_xlsx_sample_row(
            b'<row xmlns:x14ac="urn:test" r="1" x14ac:dyDescent="0.25"><c r="A1" t="inlineStr">'
            b"<is><t>Header</t></is></c></row>"
        )
        self.assertEqual("1", row_number)
        self.assertEqual("Header", values[0]["value"])

        pptx = io.BytesIO()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Presentation requirement"
        presentation.save(pptx)
        result = analyzer._extract_bytes(
            "requirement.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            pptx.getvalue(),
        )
        self.assertEqual("manual-review-required", result["status"])
        self.assertEqual(1, len(result["review_targets"]))
        self.assertIn("Presentation requirement", analyzer._collect_extracted_text(result))

    def test_xlsx_retains_late_rows_namespaces_positions_and_formulas(self):
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.title = "Controls"
        for row in range(1, 8):
            sheet.append(["HR system", "No", "No", None, "Yes"])
        sheet["A7"] = "Finance approval is mandatory."
        sheet["F7"] = "=1+1"
        sheet.merge_cells("A9:B9")
        sheet["A9"] = "Merged header"
        stream = io.BytesIO()
        book.save(stream)
        # Prefix every SpreadsheetML element, retaining namespace declarations.
        rewritten = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as original:
            with zipfile.ZipFile(rewritten, "w") as archive:
                for info in original.infolist():
                    payload = original.read(info.filename)
                    if info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml"):
                        root = analyzer.ET.fromstring(payload)
                        analyzer.ET.register_namespace("x", "http://schemas.openxmlformats.org/spreadsheetml/2006/main")
                        payload = analyzer.ET.tostring(root)
                        self.assertIn(b"<x:row", payload)
                    archive.writestr(info, payload)
        result = analyzer._extract_bytes("requirements.xlsx", "application/octet-stream", rewritten.getvalue())
        self.assertEqual("complete", result["status"])
        self.assertTrue(result["metadata"]["all_cell_evidence_retained"])
        self.assertEqual(8, result["metadata"]["rows_scanned"])
        self.assertEqual(1, result["metadata"]["formula_cells"])
        late = next(unit for unit in result["content_units"] if unit["locator"] == "sheet 'Controls' row 7")
        self.assertIn("Finance approval is mandatory.", late["text"])
        self.assertIn("B7=No | C7=No | E7=Yes", late["text"])
        self.assertEqual([1, 2, 3, 5, 6], [cell["column"] for cell in late["cells"]])
        self.assertEqual("=1+1".lstrip("="), late["cells"][-1]["formula"])
        locators, _, _ = analyzer._build_evidence_indexes({"SRC-000000000000": result})
        self.assertIn("sheet 'Controls' cell A7", locators["SRC-000000000000"])
        self.assertIn("merged_ranges=A9:B9", analyzer._collect_extracted_text(result))

    def test_xlsx_limit_is_incomplete_but_profiles_all_rows(self):
        from openpyxl import Workbook

        book = Workbook()
        for index in range(100):
            book.active.append([f"Requirement {index}", "No", "No"])
        stream = io.BytesIO()
        book.save(stream)
        with patch.object(analyzer, "MAX_EXTRACTED_TEXT_CHARS", 1500):
            result = analyzer._extract_xlsx(stream.getvalue(), analyzer.ExtractionBudget())
        self.assertEqual("manual-review-required", result["status"])
        self.assertLess(result["coverage_percent"], 100)
        self.assertEqual(100, result["metadata"]["rows_scanned"])
        self.assertGreater(result["metadata"]["rows_omitted"], 0)
        self.assertFalse(result["metadata"]["all_cell_evidence_retained"])
        self.assertTrue(result["review_targets"])

    def test_xlsx_shared_strings_preserve_late_requirement(self):
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        data = io.BytesIO()
        with zipfile.ZipFile(data, "w") as archive:
            archive.writestr("xl/workbook.xml",
                f'<workbook xmlns="{namespace}" xmlns:r="urn:relationships"><sheets>'
                '<sheet name="Requirements" r:id="rId1"/></sheets></workbook>')
            archive.writestr("xl/_rels/workbook.xml.rels",
                '<Relationships><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
            archive.writestr("xl/sharedStrings.xml",
                f'<x:sst xmlns:x="{namespace}">'
                + "".join(f"<x:si><x:t>Operational row {index}</x:t></x:si>" for index in range(6))
                + "<x:si><x:r><x:t>Finance approval </x:t></x:r>"
                  "<x:r><x:t>is mandatory.</x:t></x:r></x:si></x:sst>")
            archive.writestr("xl/worksheets/sheet1.xml",
                f'<x:worksheet xmlns:x="{namespace}"><x:sheetData>'
                + "".join(f'<x:row r="{row}"><x:c r="A{row}" t="s"><x:v>{row-1}</x:v></x:c></x:row>'
                          for row in range(1, 8))
                + "</x:sheetData></x:worksheet>")
        result = analyzer._extract_xlsx(data.getvalue(), analyzer.ExtractionBudget())
        self.assertEqual("complete", result["status"])
        row = next(unit for unit in result["content_units"] if unit["locator"] == "sheet 'Requirements' row 7")
        self.assertEqual("Finance approval is mandatory.", row["cells"][0]["value"])
        self.assertEqual("sheet 'Requirements' cell A7", row["cells"][0]["locator"])

    def test_review_batches_are_lossless_bounded_and_preserve_aliases(self):
        text = "HR system | No | No |  | Yes\n  Keep positional whitespace.  "
        self.assertEqual(text, analyzer._compact_review_text(text))
        long_text = ("Finance approval 必須 🧭\n" * 1800)
        extraction = analyzer._base_extraction("test")
        extraction["content_units"] = [
            {"locator": "row 1", "text": text},
            {"locator": "row 2", "text": text},
            {"locator": "row 3", "text": ""},
            {"locator": "row 7", "text": long_text},
        ]
        source = {"source_id": "SRC-000000000000", "sha256": "a" * 64}
        with project_directory() as directory:
            index = analyzer._write_source_review_batches(Path(directory), [source], {source["source_id"]: extraction})
            source_index = index["sources"][0]
            directory = Path(source_index["index_path"]).parent
            records = []
            for batch in source_index["batches"]:
                self.assertLessEqual(batch["byte_count"], analyzer.REVIEW_BATCH_BYTES)
                records.extend(json.loads((directory / batch["path"]).read_text(encoding="utf-8"))["records"])
            recovered = {}
            for record in records:
                recovered.setdefault(record["record_id"], "")
                recovered[record["record_id"]] += record["text"]
            self.assertEqual(text, recovered[source["source_id"] + ":unit:1"])
            self.assertEqual(long_text, recovered[source["source_id"] + ":unit:4"])
            alias = next(record for record in records if record["locator"] == "row 2")
            self.assertEqual(source["source_id"] + ":unit:1", alias["duplicate_of"])
            self.assertIn(source["source_id"] + ":unit:3", recovered)
            self.assertEqual([], analyzer._batch_integrity_errors(index, directory.parent.parent))

    def test_verified_extraction_cache_avoids_reading_source_again(self):
        with project_directory() as directory:
            run, manifest = self.prepare_fixture(Path(directory))
            source = manifest["sources"][0]
            cache_root = Path(run["output_root"]) / ".requirement-analyzer" / "cache" / f"v{analyzer.EXTRACTION_FORMAT_VERSION}"
            with patch.object(Path, "read_bytes", side_effect=AssertionError("source reread")):
                _, hit = analyzer._extract_record(source, cache_root, run["extractor_fingerprint"])
            self.assertTrue(hit)

    def test_batch_review_gate_rejects_absence_claims_before_complete_review(self):
        with project_directory() as directory:
            run, manifest = self.prepare_fixture(Path(directory))
            draft = self.make_draft_ledger(run, manifest)
            draft["batch_reviews"].pop()
            path = self.normalize_draft(run, draft, "incomplete-batches")
            failed = self.run_cli("render", "--run", str(Path(run["run_directory"]) / "run.json"),
                                  "--ledger", str(path), expected=2)
            self.assertIn("batch review is incomplete", failed["error"])
            self.assertFalse(Path(run["target_manifest_path"]).exists())

    def test_publish_shares_indexes_and_retains_independent_final_integrity(self):
        with project_directory() as directory:
            run, manifest = self.prepare_fixture(Path(directory))
            with patch.object(analyzer, "_inventory", wraps=analyzer._inventory) as inventory:
                with patch.object(analyzer, "_build_evidence_indexes", wraps=analyzer._build_evidence_indexes) as indexes:
                    published = self.publish_draft(run, manifest)
            self.assertEqual(2, inventory.call_count)
            self.assertEqual(1, indexes.call_count)
            marker = json.loads(Path(published["manifest"]).read_text(encoding="utf-8"))
            self.assertEqual("validated", marker["publication"]["status"])
            with self.assertRaisesRegex(analyzer.AnalyzerError, "already has a publication marker"):
                self.publish_draft(run, manifest)

    def test_warm_analysis_cache_precedes_model_facing_review_creation(self):
        with project_directory() as directory:
            run, manifest = self.prepare_fixture(Path(directory))
            published = self.publish_draft(run, manifest)
            with patch.object(analyzer, "_build_review_pack", side_effect=AssertionError("review pack rebuilt")):
                with patch.object(analyzer, "_write_source_review_batches", side_effect=AssertionError("batches rebuilt")):
                    prepared = self.call_in_process(analyzer._prepare, config=run["config_path"], workers=2, local_time=None)
            self.assertTrue(prepared["analysis_cache_hit"])
            Path(published["manifest"]).unlink()
            prepared = self.call_in_process(analyzer._prepare, config=run["config_path"], workers=2, local_time=None)
            self.assertFalse(prepared["analysis_cache_hit"])

    def test_final_integrity_check_blocks_source_change_during_publish(self):
        with project_directory() as directory:
            run, manifest = self.prepare_fixture(Path(directory))
            original_render = analyzer._render_markdown
            calls = 0

            def render_then_change(*arguments):
                nonlocal calls
                rendered = original_render(*arguments)
                calls += 1
                if calls == 2:
                    Path(manifest["sources"][0]["absolute_path"]).write_text(
                        "Changed after semantic validation.", encoding="utf-8"
                    )
                return rendered

            with patch.object(analyzer, "_render_markdown", side_effect=render_then_change):
                with self.assertRaisesRegex(analyzer.AnalyzerError, "Requirements files changed"):
                    self.publish_draft(run, manifest)
            self.assertFalse(Path(run["target_markdown_path"]).exists())
            self.assertFalse(Path(run["target_ledger_path"]).exists())
            self.assertFalse(Path(run["target_manifest_path"]).exists())
            self.assertTrue(analyzer._staged_paths(run)[0].exists())

    def test_batch_tampering_blocks_publication(self):
        with project_directory() as directory:
            run, manifest = self.prepare_fixture(Path(directory))
            index = json.loads(Path(run["review_index_path"]).read_text(encoding="utf-8"))
            source = index["sources"][0]
            batch_path = Path(source["index_path"]).parent / source["batches"][0]["path"]
            batch_path.write_text('{"records":[]}', encoding="utf-8")
            with self.assertRaisesRegex(analyzer.AnalyzerError, "Review batch content changed"):
                self.publish_draft(run, manifest)
            self.assertFalse(Path(run["target_manifest_path"]).exists())

    def test_navigation_tampering_blocks_publication(self):
        with project_directory() as directory:
            run, manifest = self.prepare_fixture(Path(directory))
            index = json.loads(Path(run["review_index_path"]).read_text(encoding="utf-8"))
            source = index["sources"][0]
            navigation = index["navigation"]
            targets = [
                Path(source["index_path"]).with_name("overview.json"),
                Path(source["index_path"]).with_name("index-page-000001.json"),
                Path(navigation["index_path"]).parent / navigation["batches"][0]["path"],
            ]
            for target in targets:
                with self.subTest(target=target.name):
                    original = target.read_bytes()
                    target.write_text('{"batches":[]}', encoding="utf-8")
                    with self.assertRaisesRegex(analyzer.AnalyzerError, "Review batch content changed"):
                        self.publish_draft(run, manifest)
                    self.assertFalse(Path(run["target_manifest_path"]).exists())
                    target.write_bytes(original)
            self.assertEqual([], analyzer._prepared_integrity_errors(run, manifest))

    def test_configured_project_change_blocks_old_run_publication(self):
        with project_directory() as directory:
            run, manifest = self.prepare_fixture(Path(directory))
            other = Path(directory) / "other-project"
            other.mkdir()
            config = Path(run["config_path"])
            changed = json.loads(config.read_text(encoding="utf-8"))
            changed["basePath"] = str(other.resolve())
            config.write_text(json.dumps(changed), encoding="utf-8")
            with self.assertRaisesRegex(analyzer.AnalyzerError, "basePath changed"):
                self.publish_draft(run, manifest)
            self.assertFalse(Path(run["target_manifest_path"]).exists())

    def test_quotes_are_bound_to_the_cited_locator_or_range(self):
        scoped = {}
        analyzer._build_evidence_indexes({
            "SRC-A": {"content_units": [
                {"locator": "line 1", "text": "Submit a request."},
                {"locator": "line 2", "text": "Finance approval is required."},
            ]},
        }, scoped_texts=scoped)
        texts = scoped["SRC-A"]
        self.assertTrue(analyzer._quote_at_locator("Finance approval", "line 2", texts))
        self.assertFalse(analyzer._quote_at_locator("Finance approval", "line 1", texts))
        self.assertFalse(analyzer._quote_at_locator("finance approval", "line 2", texts))
        self.assertTrue(analyzer._quote_at_locator(
            "a request. Finance approval", "lines 1-2", texts,
        ))
        self.assertFalse(analyzer._quote_at_locator(" \n ", "line 2", texts))

    def test_cell_quotes_cannot_borrow_another_cells_value(self):
        scoped = {}
        analyzer._build_evidence_indexes({
            "SRC-A": {"content_units": [{
                "locator": "sheet 'Permissions' row 1",
                "text": "A1=No | B1=Yes",
                "cells": [
                    {"locator": "sheet 'Permissions' cell A1", "value": "No", "formula": None},
                    {"locator": "sheet 'Permissions' cell B1", "value": "Yes", "formula": None},
                ],
            }]},
        }, scoped_texts=scoped)
        self.assertFalse(analyzer._quote_at_locator("Yes", "sheet 'Permissions' cell A1", scoped["SRC-A"]))
        self.assertTrue(analyzer._quote_at_locator("Yes", "sheet 'Permissions' cell B1", scoped["SRC-A"]))

    def test_source_local_complete_visual_observations_survive_one_source_change(self):
        from PIL import Image

        with project_directory() as directory:
            initial, _ = self.prepare_fixture(Path(directory))
            image_path = Path(initial["requirements_root"]) / "diagram.png"
            Image.new("RGB", (20, 20), "white").save(image_path)
            prepared = self.call_in_process(analyzer._prepare, config=initial["config_path"], workers=2, local_time=None)
            run, manifest = self.read_run(prepared)
            target = json.loads(Path(run["review_targets_path"]).read_text(encoding="utf-8"))[0]
            draft = self.make_draft_ledger(run, manifest)
            draft["manual_reviews"] = [{
                "target_id": target["target_id"], "status": "complete", "method": "complete local visual inspection",
                "coverage": target["locator"], "notes": "All visible nodes reviewed.",
                "result": "evidence-observed", "observations": [
                    {"locator": "finance node", "text": "Finance approval is mandatory."}
                ],
            }]
            path = Path(run["run_directory"]) / "reviewed.json"
            path.write_text(json.dumps(draft), encoding="utf-8")
            self.call_in_process(analyzer._publish, run=str(path.parent / "run.json"), ledger=str(path))
            brief = Path(run["requirements_root"]) / "brief.txt"
            original = brief.read_bytes()
            brief.write_bytes(original + b"\nNew unrelated requirement.")
            changed = self.call_in_process(analyzer._prepare, config=run["config_path"], workers=2, local_time=None)
            self.assertFalse(changed["analysis_cache_hit"])
            changed_draft = json.loads(Path(changed["ledger_draft"]).read_text(encoding="utf-8"))
            self.assertEqual(draft["manual_reviews"], changed_draft["manual_reviews"])
            self.assertEqual([], changed_draft["findings"])
            self.assertEqual([], changed_draft["batch_reviews"])
            changed_run, _ = self.read_run(changed)
            self.assertEqual(manifest["sources"][1]["sha256"], changed_run["reused_observations"][0]["source_sha256"])
            (Path(run["requirements_root"]) / "additional.txt").write_text(
                "A second independent corpus change.", encoding="utf-8",
            )
            multiple = self.call_in_process(analyzer._prepare, config=run["config_path"], workers=2, local_time=None)
            multiple_draft = json.loads(Path(multiple["ledger_draft"]).read_text(encoding="utf-8"))
            self.assertEqual(draft["manual_reviews"], multiple_draft["manual_reviews"])
            self.assertEqual([], multiple_draft["findings"])
            self.assertEqual([], multiple_draft["batch_reviews"])
            brief.write_bytes(original)
            Image.new("RGB", (20, 20), "black").save(image_path)
            changed_target = self.call_in_process(analyzer._prepare, config=run["config_path"], workers=2, local_time=None)
            changed_draft = json.loads(Path(changed_target["ledger_draft"]).read_text(encoding="utf-8"))
            self.assertEqual([], changed_draft["manual_reviews"])

    def test_equivalent_assertion_dedup_preserves_all_citations(self):
        with project_directory() as directory:
            run, manifest = self.prepare_fixture(Path(directory))
            draft = self.make_draft_ledger(run, manifest)
            finding = draft["findings"][0]
            duplicate = json.loads(json.dumps(finding))
            duplicate["finding_id"] = "D-999"
            duplicate["evidence"][0]["locator"] = "document metadata"
            draft["findings"].append(duplicate)
            draft["sections"]["Executive Summary"][0]["finding_ids"].append("D-999")
            normalized = json.loads(self.normalize_draft(run, draft, "dedup").read_text(encoding="utf-8"))
            merged = next(item for item in normalized["findings"] if item["statement"] == finding["statement"])
            self.assertEqual(2, len(merged["evidence"]))
            self.assertEqual(len(draft["findings"]) - 1, len(normalized["findings"]))

    def test_visual_and_unsafe_archive_require_review(self) -> None:
        from PIL import Image

        image_bytes = io.BytesIO()
        Image.new("RGB", (20, 20), "white").save(image_bytes, format="PNG")
        image = analyzer._extract_bytes("diagram.png", "image/png", image_bytes.getvalue())
        self.assertEqual("manual-review-required", image["status"])
        self.assertEqual(1, len(image["review_targets"]))

        archive_bytes = io.BytesIO()
        with zipfile.ZipFile(archive_bytes, "w") as archive:
            archive.writestr("../escape.txt", "unsafe")
        archive = analyzer._extract_bytes(
            "unsafe.zip", "application/zip", archive_bytes.getvalue()
        )
        self.assertEqual("failed", archive["status"])
        self.assertIn("Unsafe archive path", archive["warnings"][0])

        drive_archive = io.BytesIO()
        with zipfile.ZipFile(drive_archive, "w") as archive_file:
            archive_file.writestr("C:/escape.txt", "unsafe")
        archive = analyzer._extract_bytes(
            "drive.zip", "application/zip", drive_archive.getvalue()
        )
        self.assertEqual("failed", archive["status"])
        self.assertIn("Unsafe archive path", archive["warnings"][0])

        duplicate_archive = io.BytesIO()
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            with zipfile.ZipFile(duplicate_archive, "w") as archive_file:
                archive_file.writestr("duplicate.txt", "first")
                archive_file.writestr("duplicate.txt", "second")
        archive = analyzer._extract_bytes(
            "duplicate.zip", "application/zip", duplicate_archive.getvalue()
        )
        self.assertEqual("failed", archive["status"])
        self.assertIn("Duplicate archive path", archive["warnings"][0])

        from pypdf import PdfWriter

        pdf_bytes = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(pdf_bytes)
        pdf = analyzer._extract_bytes(
            "requirement.pdf", "application/pdf", pdf_bytes.getvalue()
        )
        self.assertEqual("manual-review-required", pdf["status"])
        self.assertEqual("render:page:1", pdf["review_targets"][0]["target_key"])

    def test_prepare_uses_cache_and_records_embedded_content(self) -> None:
        with project_directory() as directory:
            temporary = Path(directory)
            run, manifest = self.prepare_fixture(temporary)
            self.assertEqual(2, manifest["source_count"])
            self.assertEqual(0, run["review_target_count"])
            self.assertTrue(
                all(item["extraction_status"] == "complete" for item in manifest["sources"])
            )

            second = self.run_cli(
                "prepare",
                "--config",
                str(temporary / "source" / "lisa-config.json"),
            )
            self.assertEqual(2, second["cache_hits"])

    def test_normalize_render_and_validate_end_to_end(self) -> None:
        with project_directory() as directory:
            temporary = Path(directory)
            run, manifest = self.prepare_fixture(temporary)
            draft = self.make_draft_ledger(run, manifest)
            normalized_path = self.normalize_draft(run, draft)
            normalized = json.loads(normalized_path.read_text(encoding="utf-8"))
            self.assertTrue(
                all(
                    analyzer.FINAL_FINDING_PATTERN.fullmatch(item["finding_id"])
                    for item in normalized["findings"]
                )
            )

            rendered = self.run_cli(
                "render",
                "--run",
                str(Path(run["run_directory"]) / "run.json"),
                "--ledger",
                str(normalized_path),
            )
            markdown = Path(rendered["markdown"])
            self.assertTrue(markdown.exists())
            self.assertTrue(Path(rendered["ledger"]).exists())
            self.assertTrue(Path(rendered["manifest"]).exists())
            self.assertEqual(Path(run["run_directory"]), markdown.parent)
            self.assertFalse(Path(run["target_markdown_path"]).exists())
            self.assertFalse(Path(run["target_manifest_path"]).exists())
            rendered_run = json.loads(
                (Path(run["run_directory"]) / "run.json").read_text(encoding="utf-8")
            )
            self.assertEqual("rendered_pending_validation", rendered_run["status"])

            validated = self.run_cli(
                "validate",
                "--run",
                str(Path(run["run_directory"]) / "run.json"),
                "--ledger",
                str(normalized_path),
                "--markdown",
                str(markdown),
            )
            self.assertEqual("passed", validated["status"])
            validated_run = json.loads(
                (Path(run["run_directory"]) / "run.json").read_text(encoding="utf-8")
            )
            self.assertEqual("validated", validated_run["status"])

            cached = self.run_cli(
                "prepare",
                "--config",
                str(temporary / "source" / "lisa-config.json"),
            )
            self.assertTrue(cached["analysis_cache_hit"])
            self.assertTrue(Path(cached["reused_ledger"]).exists())
            fast = self.run_cli(
                "publish",
                "--run",
                cached["run"],
                "--ledger",
                cached["reused_ledger"],
            )
            self.assertEqual("validated", fast["status"])
            self.assertTrue(fast["analysis_cache_hit"])
            self.assertLess(fast["duration_seconds"], 30)

    def test_render_rejects_source_changes_after_prepare(self) -> None:
        with project_directory() as directory:
            temporary = Path(directory)
            run, manifest = self.prepare_fixture(temporary)
            draft = self.make_draft_ledger(run, manifest)
            normalized_path = self.normalize_draft(run, draft, "stale-source")
            (temporary / "source" / "requirements" / "brief.txt").write_text(
                "Changed requirement evidence.", encoding="utf-8"
            )
            result = self.run_cli(
                "render",
                "--run",
                str(Path(run["run_directory"]) / "run.json"),
                "--ledger",
                str(normalized_path),
                expected=2,
            )
            self.assertIn("changed after preparation", result["error"])

    def test_validation_rejects_tampered_published_artifacts(self) -> None:
        with project_directory() as directory:
            temporary = Path(directory)
            run, manifest = self.prepare_fixture(temporary)
            draft = self.make_draft_ledger(run, manifest)
            normalized_path = self.normalize_draft(run, draft, "tamper")
            rendered = self.run_cli(
                "render",
                "--run",
                str(Path(run["run_directory"]) / "run.json"),
                "--ledger",
                str(normalized_path),
            )
            markdown = Path(rendered["markdown"])
            markdown.write_text(
                markdown.read_text(encoding="utf-8") + "\nTampered prose.\n",
                encoding="utf-8",
            )
            result = self.run_cli(
                "validate",
                "--run",
                str(Path(run["run_directory"]) / "run.json"),
                "--ledger",
                str(normalized_path),
                "--markdown",
                str(markdown),
                expected=2,
            )
            self.assertIn(
                "Published Markdown differs from deterministic ledger rendering",
                result["errors"],
            )

    def test_knowledge_source_provenance_is_source_specific(self) -> None:
        with project_directory() as directory:
            temporary = Path(directory)
            run, manifest = self.prepare_fixture(temporary)
            draft = self.make_draft_ledger(run, manifest)
            source_by_name = {
                item["relative_path"]: item["source_id"] for item in manifest["sources"]
            }
            email_source = source_by_name["example.eml"]
            brief_source = source_by_name["brief.txt"]
            draft_id = f"D-{len(draft['findings']) + 1:03d}"
            draft["findings"].append(
                {
                    "finding_id": draft_id,
                    "kind": "Observed fact",
                    "statement": "The email references an external policy URL.",
                    "status": "Current",
                    "confidence": "Confirmed",
                    "evidence": [
                        {
                            "source_id": email_source,
                            "locator": "body part 1 (text/plain; utf-8)",
                            "evidence_type": "explicit",
                            "quote": "https://example.com/policy",
                        }
                    ],
                }
            )
            draft["knowledge_sources"].append(
                {
                    "name": "Policy site",
                    "classification": "Strong Candidate Knowledge Source",
                    "location": "https://example.com/policy",
                    "source_id": brief_source,
                    "hosting_type": "Public website",
                    "content_type": "Policy",
                    "structure": "Web page",
                    "grounding_purpose": "Answer policy questions",
                    "authority_priority": "Unspecified",
                    "access_behavior": "Unspecified",
                    "ownership_freshness": "Unspecified",
                    "intended_use": "Runtime grounding",
                    "finding_ids": [draft_id],
                }
            )
            normalized_path = self.normalize_draft(run, draft, "wrong-provenance")
            result = self.run_cli(
                "render",
                "--run",
                str(Path(run["run_directory"]) / "run.json"),
                "--ledger",
                str(normalized_path),
                expected=2,
            )
            self.assertIn("different source", result["error"])

    def test_schema_rejects_unknown_finding_status(self) -> None:
        with project_directory() as directory:
            temporary = Path(directory)
            run, manifest = self.prepare_fixture(temporary)
            draft = self.make_draft_ledger(run, manifest)
            draft["findings"][0]["status"] = "Unknown status"
            draft_path = Path(run["run_directory"]) / "ledger.invalid.draft.json"
            draft_path.write_text(json.dumps(draft, indent=2), encoding="utf-8")
            result = self.run_cli(
                "normalize",
                "--run",
                str(Path(run["run_directory"]) / "run.json"),
                "--ledger",
                str(draft_path),
                expected=2,
            )
            self.assertIn("Evidence ledger schema error", result["error"])

    def test_schema_requires_observed_visual_evidence_details(self) -> None:
        with project_directory() as directory:
            temporary = Path(directory)
            run, manifest = self.prepare_fixture(temporary)
            draft = self.make_draft_ledger(run, manifest)
            draft["manual_reviews"] = [
                {
                    "target_id": "REV-EXAMPLE",
                    "status": "complete",
                    "method": "visual inspection",
                    "coverage": "entire image",
                    "notes": "Reviewed.",
                    "result": "evidence-observed",
                    "observations": [],
                }
            ]
            draft_path = Path(run["run_directory"]) / "ledger.review.draft.json"
            draft_path.write_text(json.dumps(draft, indent=2), encoding="utf-8")
            result = self.run_cli(
                "normalize",
                "--run",
                str(Path(run["run_directory"]) / "run.json"),
                "--ledger",
                str(draft_path),
                expected=2,
            )
            self.assertIn("Evidence ledger schema error", result["error"])

    def test_absence_claims_require_corpus_gap_findings(self) -> None:
        with project_directory() as directory:
            temporary = Path(directory)
            run, manifest = self.prepare_fixture(temporary)
            draft = self.make_draft_ledger(run, manifest)
            draft["platform_absence_finding_ids"] = [
                draft["source_annotations"][0]["finding_ids"][0]
            ]
            normalized_path = self.normalize_draft(run, draft, "invalid-absence")
            result = self.run_cli(
                "render",
                "--run",
                str(Path(run["run_directory"]) / "run.json"),
                "--ledger",
                str(normalized_path),
                expected=2,
            )
            self.assertIn("Platform absence must cite", result["error"])

    def test_configured_knowledge_sources_are_prepopulated(self) -> None:
        with project_directory() as directory:
            base = Path(directory)
            requirements = base / "requirements"
            requirements.mkdir()
            (requirements / "brief.txt").write_text(
                "Users require a searchable assistant.", encoding="utf-8"
            )
            config = base / "lisa-config.json"
            config.write_text(
                json.dumps(
                    {
                        "basePath": ".",
                        "knowledgeSources": [
                            {
                                "name": "Procurement Policy",
                                "path": "https://contoso.sharepoint.com/policy",
                            }
                        ],
                    }
                ),
                encoding="utf-8",
            )
            prepared = self.run_cli("prepare", "--config", str(config))
            draft = json.loads(
                Path(prepared["ledger_draft"]).read_text(encoding="utf-8")
            )
            self.assertEqual(
                "Configured Knowledge Source",
                draft["knowledge_sources"][0]["classification"],
            )
            self.assertEqual("CONFIG", draft["knowledge_sources"][0]["source_id"])
            self.assertEqual(
                "https://contoso.sharepoint.com/policy",
                draft["knowledge_sources"][0]["location"],
            )

    def test_absolute_base_path_is_supported(self) -> None:
        with project_directory() as directory:
            base = Path(directory)
            requirements = base / "requirements"
            requirements.mkdir()
            (requirements / "brief.txt").write_text("Requirement", encoding="utf-8")
            config = base / "lisa-config.json"
            config.write_text(
                json.dumps({"basePath": str(base.resolve())}),
                encoding="utf-8",
            )
            result = self.run_cli(
                "prepare",
                "--config",
                str(config),
            )
            self.assertEqual("prepared", result["status"])
            self.assertTrue((base / "output" / "analysis").is_dir())

    def test_default_temp_output_uses_analysis_child(self) -> None:
        with project_directory() as directory:
            base = Path(directory)
            requirements = base / "requirements"
            requirements.mkdir()
            (requirements / "brief.txt").write_text(
                "Users require a searchable assistant.", encoding="utf-8"
            )
            config = base / "lisa-config.json"
            config.write_text(json.dumps({"basePath": "."}), encoding="utf-8")
            prepared = self.run_cli(
                "prepare",
                "--config",
                str(config),
            )
            run = json.loads(Path(prepared["run"]).read_text(encoding="utf-8"))
            self.assertEqual((base / "output").resolve(), Path(run["temp_output_path"]))
            self.assertEqual(
                (base / "output" / "analysis").resolve(),
                Path(run["output_root"]),
            )
            self.assertEqual(
                ["analysis"],
                sorted(item.name for item in (base / "output").iterdir()),
            )


if __name__ == "__main__":
    unittest.main()
